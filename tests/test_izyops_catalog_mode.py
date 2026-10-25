"""Who manages the products (IZY Ops / website / both), and the backups that
keep the website's products safe."""
import base64
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Product, ProductBackup, ProductImage, ProductVariant
from app.services import backups


@pytest.fixture
def api(app):
    app.config.update(IZYOPS_API_USERNAME='izy', IZYOPS_API_PASSWORD='secret')
    token = base64.b64encode(b'izy:secret').decode()
    client = app.test_client()

    def call(method, url, **kw):
        return getattr(client, method)(url, headers={'Authorization': f'Basic {token}'}, **kw)
    return call


def _mode(api, mode):
    return api('post', '/api/izyops/settings', json={'catalog_managed_by': mode})


@pytest.fixture
def tee(db, category):
    p = Product(name='Basic Tee', slug='basic-tee', sku='TEE', price=Decimal('9.99'), stock_quantity=20,
                description='Soft cotton tee', short_description='Soft', category=category, is_active=True)
    db.session.add(p)
    db.session.flush()
    db.session.add(ProductImage(product_id=p.id, image_url='https://img.test/tee.jpg', is_primary=True))
    db.session.commit()
    return p


# ─── the setting ─────────────────────────────────────────────────────────────

def test_settings_store_the_mode_and_back_up_everything_once(api, db, tee):
    data = _mode(api, 'both').get_json()['data']
    assert data == {'catalog_managed_by': 'both', 'forced_by_server': False}
    first = ProductBackup.query.filter_by(batch=backups.FIRST_BATCH).all()
    assert len(first) == 1 and backups.data_of(first[0])['images'] == ['https://img.test/tee.jpg']
    _mode(api, 'izyops')
    assert ProductBackup.query.filter_by(batch=backups.FIRST_BATCH).count() == 1
    assert _mode(api, 'nope').status_code == 400


def test_lock_follows_the_mode(app, api, admin_client, db, tee):
    url = f'/admin/products/{tee.id}/toggle'
    _mode(api, 'izyops')
    assert admin_client.post(url).status_code == 302 and db.session.get(Product, tee.id).is_active
    assert 'Products are managed in IZY Ops' in admin_client.get('/admin/products').get_data(as_text=True)
    _mode(api, 'both')
    assert admin_client.post(url).status_code == 200 and not db.session.get(Product, tee.id).is_active
    assert 'Edit here or in IZY Ops' in admin_client.get('/admin/products').get_data(as_text=True)
    app.config['IZYOPS_OWNS_CATALOG'] = True  # the server setting still forces the lock
    assert admin_client.post(url).status_code == 302


def test_nothing_is_locked_before_izyops_says_so(admin_client, db, tee):
    assert admin_client.post(f'/admin/products/{tee.id}/toggle').status_code == 200


def test_website_only_refuses_product_pushes(api, db, tee):
    _mode(api, 'website')
    resp = api('put', '/api/izyops/products/TEE', json={'price': '1.00'})
    assert resp.status_code == 409 and resp.get_json()['error'] == 'managed_on_website'
    assert db.session.get(Product, tee.id).price == Decimal('9.99')


# ─── nothing lost ────────────────────────────────────────────────────────────

def test_a_push_saves_the_product_first_once_an_hour(api, db, tee):
    _mode(api, 'both')
    api('put', '/api/izyops/products/TEE', json={'price': '11.00'})
    api('put', '/api/izyops/products/TEE', json={'price': '12.00'})
    saved = ProductBackup.query.filter_by(product_id=tee.id, batch=None).all()
    assert len(saved) == 1 and backups.data_of(saved[0])['price'] == '9.99'
    saved[0].taken_at = datetime.utcnow() - timedelta(hours=2)
    db.session.commit()
    api('put', '/api/izyops/products/TEE', json={'price': '12.00'})  # no change: no backup
    assert ProductBackup.query.filter_by(product_id=tee.id, batch=None).count() == 1
    api('put', '/api/izyops/products/TEE', json={'price': '13.00'})
    assert ProductBackup.query.filter_by(product_id=tee.id, batch=None).count() == 2


def test_an_empty_description_from_izyops_never_erases_the_websites(api, db, tee):
    _mode(api, 'izyops')
    api('put', '/api/izyops/products/TEE', json={'description': None, 'short_description': '', 'name': 'Tee'})
    product = db.session.get(Product, tee.id)
    assert product.name == 'Tee' and product.description == 'Soft cotton tee' and product.short_description == 'Soft'


def test_restore_puts_the_old_values_and_photos_back(api, admin_client, db, tee):
    _mode(api, 'both')
    api('put', '/api/izyops/products/TEE', json={'name': 'Tee v2', 'price': '19.99',
                                                  'images': ['https://izy.test/new.jpg']})
    product = db.session.get(Product, tee.id)
    assert product.name == 'Tee v2' and [i.image_url for i in product.gallery_images] == ['https://izy.test/new.jpg']
    backup = ProductBackup.query.filter_by(product_id=tee.id, batch=None).one()
    page = admin_client.get('/admin/products/backups').get_data(as_text=True)
    assert 'Basic Tee' in page and 'Before an IZY Ops change' in page
    assert admin_client.post(f'/admin/products/backups/{backup.id}/restore').status_code == 302
    product = db.session.get(Product, tee.id)
    assert product.name == 'Basic Tee' and product.price == Decimal('9.99')
    assert [i.image_url for i in product.gallery_images] == ['https://img.test/tee.jpg']
    assert ProductBackup.query.filter_by(product_id=tee.id, batch=None).count() == 2  # the undone values too


def test_restore_is_off_while_izyops_manages_products(api, admin_client, db, tee):
    _mode(api, 'izyops')
    api('put', '/api/izyops/products/TEE', json={'name': 'Tee v2'})
    backup = ProductBackup.query.filter_by(product_id=tee.id, batch=None).one()
    admin_client.post(f'/admin/products/backups/{backup.id}/restore')
    assert db.session.get(Product, tee.id).name == 'Tee v2'


def test_the_download_has_photos_and_restores_them(admin_client, db, tee):
    resp = admin_client.get('/admin/products/export')
    ws = load_workbook(BytesIO(resp.data)).active
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row['image_urls'] == 'https://img.test/tee.jpg' and row['product_name'] == 'Basic Tee'

    product = db.session.get(Product, tee.id)
    product.images.clear()
    db.session.commit()
    from app.services.product_import import apply_rows, parse_workbook
    rows, errors = parse_workbook(resp.data)
    result = apply_rows(rows)
    assert not errors and result['updated'] == 1
    assert [i.image_url for i in db.session.get(Product, tee.id).gallery_images] == ['https://img.test/tee.jpg']


def test_the_first_backup_downloads_as_excel(api, admin_client, db, tee, variant_product):
    _mode(api, 'both')
    resp = admin_client.get('/admin/products/backups/download')
    ws = load_workbook(BytesIO(resp.data)).active
    names = [r[1] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert names.count('Basic Tee') == 1 and names.count('Color Mug') == 3


def test_a_colour_change_marks_the_product_changed(db, variant_product):
    mug, red = variant_product[0], variant_product[1]
    mug.updated_at = datetime(2026, 1, 1)
    db.session.commit()
    db.session.get(ProductVariant, red.id).price_override = Decimal('8.00')
    db.session.commit()
    assert db.session.get(Product, mug.id).updated_at > datetime(2026, 6, 1)
